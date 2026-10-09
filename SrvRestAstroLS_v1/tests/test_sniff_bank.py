from __future__ import annotations

from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from services.ingest.sniff_bank import sniff_file
from routes.v1 import reconcile_start


def _make_extract_with_bad_dimension(
    path: Path,
    *,
    account: str = "100-393300535-000",
) -> None:
    """Genera un XLSX válido cuyo XML declara erróneamente ``A1`` como dimensión."""
    wb = Workbook()
    ws = wb.active
    ws.title = "principal"
    ws.append(["Extracto de Cuenta"])
    ws.append(["Tipo y Nro. de Cuenta", f"CC $ {account}"])
    ws.append(["Fecha desde", "14/07/2026"])
    ws.append(["Fecha hasta", "14/07/2026"])
    ws.append([])
    ws.append(["Concepto/Cod.Op.", "Fecha", "Comprobante", "Sucursal", "Importe", "Descripción", "Saldo"])
    ws.append(["ORD.PGO.ME", "14/07/2026", "123", "010", -100.5, "Pago", 999.5])
    ws.append([None, None, None, None, None, None, "Saldo Final: 999.5"])
    wb.save(path)

    with ZipFile(path) as source:
        contents = {item.filename: source.read(item.filename) for item in source.infolist()}

    sheet_xml = contents["xl/worksheets/sheet1.xml"]
    contents["xl/worksheets/sheet1.xml"] = sheet_xml.replace(
        b'<dimension ref="A1:G8"/>', b'<dimension ref="A1"/>', 1
    )

    with ZipFile(path, "w", compression=ZIP_DEFLATED) as target:
        for name, content in contents.items():
            target.writestr(name, content)


def test_sniff_bank_extract_ignores_incorrect_worksheet_dimension(tmp_path: Path) -> None:
    path = tmp_path / "extracto_dimension_incorrecta.xlsx"
    _make_extract_with_bad_dimension(path)

    result = sniff_file(path)

    assert result["kind"] == "bank_movements"
    assert result["detected"]["bank"] == "patagonia"
    assert result["validation"]["is_valid"] is True
    assert result["validation"]["header_row"] == 6
    assert result["validation"]["rows_count"] == 1


def test_sniff_detects_confirmed_banco_nacion_account(tmp_path: Path) -> None:
    path = tmp_path / "extracto_bna.xlsx"
    _make_extract_with_bad_dimension(path, account="00.060.678/47")

    result = sniff_file(path)

    assert result["kind"] == "bank_movements"
    assert result["detected"]["bank"] == "nacion"
    assert result["detected"]["account_full"] == "00.060.678/47"
    assert result["needs"]["bank"] is False


def test_sniff_detects_additional_client_accounts_with_short_or_full_format(tmp_path: Path) -> None:
    cases = [
        ("64464/16", "nacion", "00.064.464/16"),
        ("00.300.318/81", "nacion", "00.300.318/81"),
        ("300011/39", "nacion", "00.300.011/39"),
        ("300285/28", "nacion", "00.300.285/28"),
        ("300342/90", "nacion", "00.300.342/90"),
        ("1914/85", "nacion", "00.001.914/85"),
        ("3050233/8", "nacion", "03.050.233/8"),
        ("50289/6", "provincia", "50289/6"),
        ("260055", "ciudad", "3-111-0100026005-5"),
    ]
    for index, (account, bank, expected_full) in enumerate(cases):
        path = tmp_path / f"extracto_cuenta_{index}.xlsx"
        _make_extract_with_bad_dimension(path, account=account)

        result = sniff_file(path)

        assert result["detected"]["bank"] == bank
        assert result["detected"]["account_full"] == expected_full
        assert result["needs"]["bank"] is False


def test_pilaga_saldos_ignores_empty_rows_from_read_only_workbook(monkeypatch, tmp_path: Path) -> None:
    class Worksheet:
        def iter_rows(self, values_only: bool = False):
            return iter([(), ("Saldo Inicial: 10.00",), (), ("Saldo Final: 25.50",)])

    class WorkbookStub:
        sheetnames = ["Resumen cuenta bancaria"]

        def __getitem__(self, name: str):
            return Worksheet()

        def close(self) -> None:
            pass

    monkeypatch.setattr(reconcile_start, "load_workbook", lambda *args, **kwargs: WorkbookStub())

    assert reconcile_start._get_pilaga_saldos(tmp_path / "pilaga.xlsx") == (10.0, 25.5)
